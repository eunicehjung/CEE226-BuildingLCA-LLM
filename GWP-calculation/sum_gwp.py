"""
sum_gwp.py — filter full_lca_results.xlsx and sum GWP per project.

HOW TO USE
  1. Edit the SETTINGS block below.
  2. Run:  python sum_gwp.py
  3. Results print to the terminal and are saved to OUTPUT_FILE (CSV).
     Warnings, if any, go to a second CSV ending in _warnings.csv.
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from datetime import datetime

# =============================================================================
# SETTINGS — edit these
# =============================================================================
SCRIPT_DIR = Path(__file__).resolve().parent

BASE_DIR = SCRIPT_DIR
INPUT_FILE = BASE_DIR / "full_lca_results.xlsx"
OUTPUT_FILE = BASE_DIR / f"{datetime.now():%Y%m%d_%H}_gwp-total.csv"


# 1) Projects: a list of numbers, e.g. [1, 5, 12], or "all"
PROJECTS = [1] #"all"

# 2) Elements to include (comment out a line with # to exclude it)
ELEMENTS = [
    "Substructure",
    "Shell - Superstructure",
    "Shell - Enclosure",
    "Interiors - Construction",
    "Interiors - Finishes",
    # "Unknown",
]

# 3) Life cycle stages to include (comment out a line with # to exclude it)
STAGES = [
    "A1-A3",
    "A4",
    "B4-B5",
    "C2-C4",
    # "D",
]
# =============================================================================


def load_data(path: Path) -> pd.DataFrame:
    """Read the Excel file. ~0.9M rows is slow to parse, so a .pkl cache is
    written next to it and reused as long as the Excel file hasn't changed."""
    cache = path.with_suffix(".pkl")
    if cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
        return pd.read_pickle(cache)
    print(f"Reading {path.name} (first run is slow; caching to {cache.name}) ...")
    df = pd.read_excel(path, usecols=["project_index", "omniclass_element",
                                      "life_cycle_stage", "gwp"])
    df.to_pickle(cache)
    return df


def validate_settings(df):
    """Stop early on typos, since a misspelled element silently matches nothing."""
    for name, chosen, col in [("ELEMENTS", ELEMENTS, "omniclass_element"),
                              ("STAGES", STAGES, "life_cycle_stage")]:
        valid = set(df[col].dropna().unique())
        bad = [v for v in chosen if v not in valid]
        if bad:
            raise SystemExit(f"{name} contains values not in the file: {bad}\n"
                             f"Valid options: {sorted(valid)}")


def find_gaps(df, projects):
    """List requested projects missing a requested element or stage, so a
    smaller-than-expected total is explained."""
    rows = []
    sub = df[df["project_index"].isin(projects)]
    for p in projects:
        g = sub[sub["project_index"] == p]
        if g.empty:
            rows.append((p, "project not found in file"))
            continue
        for col, vals in [("omniclass_element", ELEMENTS), ("life_cycle_stage", STAGES)]:
            missing = [v for v in vals if v not in set(g[col])]
            if missing:
                rows.append((p, f"no rows for {col} = {', '.join(missing)}"))
    return pd.DataFrame(rows, columns=["project_index", "issue"])


# def write_excel(totals, gaps, path):
#     settings = pd.DataFrame({
#         "setting": ["input_file", "projects", "elements", "life_cycle_stages", "units"],
#         "value": [INPUT_FILE, str(PROJECTS), ", ".join(ELEMENTS), ", ".join(STAGES),
#                   "kg CO2e"],
#     })
#     with pd.ExcelWriter(path, engine="openpyxl") as xw:
#         totals.to_excel(xw, sheet_name="gwp_totals", index=False)
#         settings.to_excel(xw, sheet_name="settings", index=False)
#         gaps.to_excel(xw, sheet_name="warnings", index=False)
#
#     # Light formatting: Arial, thousands separators, readable column widths
#     wb = load_workbook(path)
#     for ws in wb.worksheets:
#         for row in ws.iter_rows():
#             for c in row:
#                 c.font = Font(name="Arial", bold=(c.row == 1))
#         for col in ws.columns:
#             width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
#             ws.column_dimensions[col[0].column_letter].width = min(width + 2, 80)
#     for (cell,) in wb["gwp_totals"].iter_rows(min_row=2, min_col=2, max_col=2):
#         cell.number_format = "#,##0.0"
#     wb.save(path)


def main():
    df = load_data(Path(INPUT_FILE))
    validate_settings(df)

    projects = sorted(df["project_index"].unique()) if PROJECTS == "all" else list(PROJECTS)

    sel = df[df["project_index"].isin(projects)
             & df["omniclass_element"].isin(ELEMENTS)
             & df["life_cycle_stage"].isin(STAGES)]

    # reindex keeps every requested project; ones with no matching rows get 0
    totals = (sel.groupby("project_index")["gwp"].sum()
                 .reindex(projects, fill_value=0.0)
                 .rename("total_gwp_kgCO2e")
                 .reset_index())
    gaps = find_gaps(df, projects)

    print(f"Elements: {ELEMENTS}\nStages:   {STAGES}")
    print(f"Rows used: {len(sel):,}\n")
    print(totals.to_string(index=False, float_format=lambda x: f"{x:,.1f}"))
    if len(gaps):
        warn_file = OUTPUT_FILE.with_name(OUTPUT_FILE.stem + "_warnings.csv")
        gaps.to_csv(warn_file, index=False)
        print(f"\n{len(gaps)} warning(s) — saved to {warn_file.name}")

    totals.to_csv(OUTPUT_FILE, index=False)
    print(f"\nSaved to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
